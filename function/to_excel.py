import openpyxl
from openpyxl import worksheet

COLUMN_EXCEL = ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 'N', 'O', 'P', 'Q', 'R', 'S', 'T', 'U', 'V', 'W', 'X', 'Y', 'Z']


def table_to_excel(table, path):
    wb = openpyxl.Workbook()
    ws = wb.active

    for col in range(table.columnCount()):
        head = table.horizontalHeaderItem(col).text()
        ws["%s%s" % (COLUMN_EXCEL[col], 1)] = str(head)

        for row in range(table.rowCount()):
            for col in range(table.columnCount()):
                if table.item(row, col):
                    try:
                        txt = float(table.item(row, col).text().replace(" ", ""))
                    except:
                        txt = table.item(row, col).text()
                else:
                    txt = ''
                ws["%s%s" % (COLUMN_EXCEL[col], row+2)] = txt

    wb.save(path + ".xlsx")


def patch_worksheet():
    def merge_cells(self, range_string=None, start_row=None, start_column=None, end_row=None, end_column=None):
        """ Set merge on a cell range.  Range is a cell range (e.g. A1:E1)
        This is monkeypatched to remove cell deletion bug
        https://bitbucket.org/openpyxl/openpyxl/issues/365/styling-merged-cells-isnt-working
        """
        if not range_string and not all((start_row, start_column, end_row, end_column)):
            msg = "You have to provide a value either for 'coordinate' or for\
            'start_row', 'start_column', 'end_row' *and* 'end_column'"
            raise ValueError(msg)
        elif not range_string:
            range_string = '%s%s:%s%s' % (get_column_letter(start_column),
                                          start_row,
                                          get_column_letter(end_column),
                                          end_row)
        elif ":" not in range_string:
            if COORD_RE.match(range_string):
                return  # Single cell, do nothing
            raise ValueError("Range must be a cell range (e.g. A1:E1)")
        else:
            range_string = range_string.replace('$', '')

        if range_string not in self._merged_cells:
            self._merged_cells.append(range_string)


        # The following is removed by this monkeypatch:

        # min_col, min_row, max_col, max_row = range_boundaries(range_string)
        # rows = range(min_row, max_row+1)
        # cols = range(min_col, max_col+1)
        # cells = product(rows, cols)

        # all but the top-left cell are removed
        #for c in islice(cells, 1, None):
            #if c in self._cells:
                #del self._cells[c]

    worksheet.Worksheet.merge_cells = merge_cells